import re, zipfile
from collections import Counter
from openpyxl import load_workbook
from .security import phone, keyed
from .models import Member

def parse(upload):
    if upload.size > 5*1024*1024 or not upload.name.lower().endswith('.xlsx'): raise ValueError('5MB 이하 XLSX만 가능합니다.')
    with zipfile.ZipFile(upload) as archive:
        if sum(x.file_size for x in archive.infolist()) > 30*1024*1024: raise ValueError('압축 해제 크기가 너무 큽니다.')
    upload.seek(0)
    book=load_workbook(upload,read_only=True,data_only=False)
    try:
        if '교인명부' not in book.sheetnames: raise ValueError('교인명부 시트가 필요합니다.')
        sheet=book['교인명부']
        if sheet['B2'].value != '이름' or sheet['C2'].value != '전화번호': raise ValueError('2행 B열 이름, C열 전화번호 헤더가 필요합니다.')
        rows=[]
        for index, cells in enumerate(sheet.iter_rows(min_row=3,max_col=3),3):
            if index > 10003: raise ValueError('최대 10,000행입니다.')
            namecell,numcell=cells[1:3]; name=namecell.value; number=numcell.value
            if name in (None,'') and number in (None,''): continue
            name=str(name or '').strip(); number=phone(number or '')
            issue=''
            if not name or not number: issue='이름·전화번호 누락'
            elif namecell.data_type == 'f' or numcell.data_type == 'f': issue='수식 확인 필요'
            elif len(name)>80: issue='이름 길이 확인 필요'
            elif re.fullmatch(r'1\d{8,9}',number): issue='앞자리 0 복구 후보: 0'+number
            elif not re.fullmatch(r'0\d{8,10}',number): issue='번호 확인 필요'
            rows.append({'row':index,'name':name,'phone':number,'issue':issue})
        pairs=Counter((r['name'],r['phone']) for r in rows)
        names=Counter(r['name'] for r in rows); phones=Counter(r['phone'] for r in rows)
        existing=list(Member.objects.all())
        for r in rows:
            matches=[m for m in existing if m.name == r['name'] and m.phone == r['phone']]
            r['existing']=str(matches[0].id) if len(matches)==1 else ''
            r['baseline']={str(m.id):keyed(str(m.id)+':'+str(m.updated_at)) for m in existing if m.name==r['name'] or m.phone==r['phone']}
            r['ambiguous']=len(matches)>1
            r['duplicate']=pairs[r['name'],r['phone']]>1
            r['same_name']=names[r['name']]>1; r['shared_phone']=phones[r['phone']]>1
            r['change_candidates']=[{'id':str(m.id),'name':m.name} for m in existing if not matches and (m.name==r['name'] or m.phone==r['phone'])]
        present={(r['name'],r['phone']) for r in rows}
        summary={'valid':sum(not r['issue'] and not r['duplicate'] and not r['ambiguous'] for r in rows),'missing_or_invalid':sum(bool(r['issue']) for r in rows),'duplicates':sum(r['duplicate'] for r in rows),'same_names':sum(r['same_name'] for r in rows),'shared_phones':sum(r['shared_phone'] for r in rows),'added':sum(not r['existing'] for r in rows),'change_candidates':sum(bool(r['change_candidates']) for r in rows),'omitted':[{'id':str(m.id),'name':m.name} for m in existing if (m.name,m.phone) not in present]}
        return rows,summary
    finally: book.close()
