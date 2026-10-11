import io
import json
import secrets
import zipfile
from collections import Counter
from openpyxl import load_workbook
from flask import abort, g
from .db import db, now, audit
from .security import decrypt, encrypt, digest, normalize_phone, valid_phone


def preview(raw):
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 25_000_000 or len(archive.infolist())>1000:
                abort(400, '압축 해제 크기 제한을 초과했습니다.')
        book=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
        if '교인명부' not in book.sheetnames:
            abort(400,'교인명부 시트가 필요합니다.')
        sheet=book['교인명부']
        rows=sheet.iter_rows(values_only=True,max_col=3)
        next(rows,None)
        header=next(rows,())
        if len(header)<3 or header[1]!='이름' or header[2]!='전화번호':
            abort(400,'2행 B열 이름 / C열 전화번호를 확인하세요.')
        parsed=[]
        for rowno,row in enumerate(rows,3):
            name=str(row[1] or '').strip() if len(row)>1 else ''
            value=row[2] if len(row)>2 else None
            if not name and value in (None,''):
                continue
            if isinstance(value,float) and value.is_integer():
                value=int(value)
            if len(parsed)>=10000:
                abort(400,'최대 10,000명의 명부를 처리합니다.')
            if len(name)>100 or len(str(value or ''))>100:
                abort(400,'이름·전화번호 길이를 확인하세요.')
            phone=normalize_phone(value or '')
            issue=''
            candidate=''
            if phone and not phone.startswith('0'):
                candidate='0'+phone
                issue='앞자리 0 복구 후보: 관리자 확인 필요'
            if not name or not phone:
                issue='이름 또는 전화번호 누락'
            elif not valid_phone(phone) and not issue:
                issue='전화번호 형식 확인 필요'
            parsed.append({'row':rowno,'name':name,'phone':phone,'candidate':candidate,'issue':issue})
        book.close()
    except (zipfile.BadZipFile,KeyError,ValueError,TypeError):
        abort(400,'유효한 XLSX 파일이 아닙니다.')
    old=[{'id':x['id'],'name':decrypt(x['name']),'phone':decrypt(x['phone']),'active':bool(x['active']),'version':x['version']} for x in db().execute('SELECT * FROM members')]
    pairs=Counter((x['name'],x['phone']) for x in parsed)
    names=Counter(x['name'] for x in parsed)
    phones=Counter(x['phone'] for x in parsed if x['phone'])
    for row in parsed:
        exact=[x for x in old if x['name']==row['name'] and x['phone']==row['phone']]
        row['duplicate']=pairs[(row['name'],row['phone'])]>1
        row['shared_phone']=phones[row['phone']]>1
        row['same_name']=names[row['name']]>1
        row['existing_id']=exact[0]['id'] if len(exact)==1 else None
        row['possible_changes']=[x['id'] for x in old if x['name']==row['name'] or x['phone']==row['phone']] if not exact else []
    existing={r['existing_id'] for r in parsed}
    payload={'rows':parsed,'existing':old,'missing':[x['id'] for x in old if x['id'] not in existing], 'valid_count':sum(not x['issue'] and not x['duplicate'] for x in parsed)}
    token=secrets.token_urlsafe(24)
    db().execute('DELETE FROM imports WHERE expires<=?',(now(),))
    db().execute('INSERT INTO imports VALUES(?,?,?,?)',(token,g.identity['id'],encrypt(json.dumps(payload,ensure_ascii=False)),now()+1800))
    return {'import_id':token,**payload}


def apply(import_id, decisions):
    connection=db()
    connection.execute('BEGIN IMMEDIATE')
    try:
        item=connection.execute('SELECT * FROM imports WHERE id=? AND admin_id=? AND expires>?',(import_id,g.identity['id'],now())).fetchone()
        if not item:
            abort(400,'미리보기가 만료되었습니다.')
        payload=json.loads(decrypt(item['payload']))
        original={x['id']:x for x in payload['existing']}
        current={x['id']:x for x in connection.execute('SELECT * FROM members')}
        if set(original)!=set(current) or any(original[k]['version']!=current[k]['version'] for k in original):
            abort(409,'명부가 변경되었습니다. 다시 업로드하세요.')
        touched=set()
        pairset=set()
        for row in payload['rows']:
            decision=decisions.get(str(row['row']),{})
            action=decision.get('action','keep' if row['existing_id'] else 'skip')
            if action=='skip':
                continue
            if action not in ('keep','create','link'):
                abort(400)
            name=str(decision.get('name',row['name'])).strip()
            phone=normalize_phone(decision.get('phone',row['phone']))
            if not name or len(name)>100 or not valid_phone(phone):
                abort(400,'이름과 번호를 확인하세요.')
            if (name,phone) in pairset:
                abort(400,'동일한 이름·번호를 중복 반영할 수 없습니다.')
            pairset.add((name,phone))
            target=row['existing_id'] if action=='keep' else decision.get('member_id') if action=='link' else None
            if target:
                if target not in current or target in touched:
                    abort(400,'중복되거나 유효하지 않은 ID입니다.')
                if action=='link' and not decision.get('confirmed'):
                    abort(400,'기존 교인 신원 확인이 필요합니다.')
                old=current[target]
                if decrypt(old['name'])!=name or decrypt(old['phone'])!=phone:
                    connection.execute('UPDATE members SET name=?,phone=?,phone_key=?,version=version+1,updated=? WHERE id=?',(encrypt(name),encrypt(phone),digest(phone),now(),target))
                    connection.execute('DELETE FROM sessions WHERE member_id=?',(target,))
                    connection.execute('UPDATE challenges SET used=1 WHERE member_id=?',(target,))
                    audit(g.identity['id'],'member_updated',target,'명부 확인 반영')
                touched.add(target)
            else:
                if action!='create' or not decision.get('confirmed'):
                    abort(400,'신규 등록 확인이 필요합니다.')
                if any(decrypt(x['name'])==name and decrypt(x['phone'])==phone for x in current.values()):
                    abort(400,'기존 교인이 있습니다. 기존 ID를 연결하세요.')
                target=secrets.token_hex(16)
                connection.execute('INSERT INTO members VALUES(?,?,?,?,1,1,?,?)',(target,encrypt(name),encrypt(phone),digest(phone),now(),now()))
                audit(g.identity['id'],'member_created',target)
        # Missing entries deliberately stay active; separate reasoned action required.
        connection.execute('DELETE FROM imports WHERE id=?',(import_id,))
        audit(g.identity['id'],'roster_import',import_id)
        connection.execute('COMMIT')
    except BaseException:
        connection.execute('ROLLBACK')
        raise
    return {'ok':True}
